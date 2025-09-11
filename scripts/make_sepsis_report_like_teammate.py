# -*- coding: utf-8 -*-
"""
Create a paper-style Word report for the SEPSIS final run, mimicking the teammate's report.
- Source folder (fixed): reports/final_thr28
- Inputs: metrics.json, roc.png, pr.png, confusion_matrix.png, panel_usage.png, case_studies.txt (optional)
- Output: reports/final_thr28/Sepsis_Final_Report.docx

No comparisons. No calibration curves. Uses panel names (CBC, CMP, aPTT, ABG).
"""

import os, json, math, datetime
from pathlib import Path

# --- 3rd party ---
# pip install python-docx pillow
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from PIL import Image

# ------------------------------- CONFIG -------------------------------- #
RUN_DIR = Path("reports/final_thr28")  # fixed as requested
OUT_DOC = RUN_DIR / "Sepsis_Final_Report.docx"

IMG_NAMES = {
    "ROC": "roc.png",
    "PR": "pr.png",
    "Confusion": "confusion_matrix.png",
    "Panels": "panel_usage.png",
}
METRICS_JSON = "metrics.json"
CASE_STUDIES = "case_studies.txt"

# If your figures are very big, they'll be resized to fit page width:
MAX_IMAGE_WIDTH_IN = 6.2  # inches for A4/Letter margins
# ----------------------------------------------------------------------- #


def _safe_load_json(path: Path, fallback: dict = None):
    try:
        with open(path, "r", encoding="utf-8") as f:
            return json.load(f)
    except Exception as e:
        print(f"[warn] Could not load JSON: {path} -> {e}")
        return fallback or {}


def _insert_heading(doc: Document, text: str, level: int = 1):
    p = doc.add_heading(text, level=level)
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    return p


def _insert_kv_table(doc: Document, rows, col_widths_in=(2.2, 4.2), first_col_bold=True):
    """
    rows: list of (key, value) where value can be str or number
    """
    table = doc.add_table(rows=1, cols=2)
    table.style = "Light List Accent 1"
    hdr_cells = table.rows[0].cells
    hdr_cells[0].text = "Metric"
    hdr_cells[1].text = "Value"

    for k, v in rows:
        row_cells = table.add_row().cells
        row_cells[0].text = str(k)
        row_cells[1].text = str(v)

    # widths
    for i, w in enumerate(col_widths_in):
        for row in table.rows:
            row.cells[i].width = Inches(w)

    # bold first column if requested
    if first_col_bold:
        for r in table.rows[1:]:
            for run in r.cells[0].paragraphs[0].runs:
                run.bold = True

    doc.add_paragraph()  # spacing


def _fmt_pct(x):
    try:
        return f"{100.0 * float(x):.2f}%"
    except Exception:
        return "—"


def _fmt_money(x):
    try:
        return f"${float(x):.2f}"
    except Exception:
        return "—"


def _fmt_float(x, nd=3):
    try:
        return f"{float(x):.{nd}f}"
    except Exception:
        return "—"


def _compute_derived(m):
    # robust pulls with defaults
    tp = int(m.get("tp", 0))
    tn = int(m.get("tn", 0))
    fp = int(m.get("fp", 0))
    fn = int(m.get("fn", 0))
    n = int(m.get("n_episodes", tp + tn + fp + fn))

    # prevalence
    prevalence = (tp + fn) / n if n else float("nan")

    # specificity (already in JSON but recompute for safety)
    specificity = tn / (tn + fp) if (tn + fp) else float("nan")
    sensitivity = tp / (tp + fn) if (tp + fn) else float("nan")

    # NPV, PPV
    ppv = tp / (tp + fp) if (tp + fp) else float("nan")
    npv = tn / (tn + fn) if (tn + fn) else float("nan")

    # F1/$ efficiency (just an illustration, not in your teammate’s doc necessarily)
    f1 = float(m.get("f1", float("nan")))
    avg_cost = float(m.get("avg_cost", float("nan")))
    f1_per_dollar = f1 / avg_cost if (avg_cost and not math.isnan(avg_cost)) else float("nan")

    return dict(
        prevalence=prevalence,
        specificity=specificity,
        sensitivity=sensitivity,
        ppv=ppv,
        npv=npv,
        f1_per_dollar=f1_per_dollar,
    )


def _insert_image_if_exists(doc: Document, img_path: Path, caption: str):
    if not img_path.exists():
        doc.add_paragraph(f"[Missing figure: {img_path.name}]")
        return

    # resize to fit width while preserving aspect ratio
    try:
        with Image.open(img_path) as im:
            width, height = im.size
            dpi = im.info.get("dpi", (300, 300))[0]
            # convert to inches using a sane default dpi if missing
            width_in = width / (dpi if dpi else 300)
            scale = 1.0
            if width_in > MAX_IMAGE_WIDTH_IN:
                scale = MAX_IMAGE_WIDTH_IN / width_in
            doc.add_picture(str(img_path), width=Inches(width_in * scale))
    except Exception:
        # if PIL fails for any reason, let python-docx try the raw insert
        doc.add_picture(str(img_path), width=Inches(MAX_IMAGE_WIDTH_IN))

    cap = doc.add_paragraph(caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cap.runs[0].italic = True
    doc.add_paragraph()  # spacing


def _append_case_studies(doc: Document, path: Path):
    if not path.exists():
        return
    _insert_heading(doc, "Appendix: Case Studies", level=1)
    try:
        with open(path, "r", encoding="utf-8") as f:
            content = f.read().strip()
    except Exception as e:
        doc.add_paragraph(f"[Could not read case_studies.txt: {e}]")
        return

    if not content:
        doc.add_paragraph("[No case studies recorded.]")
        return

    # chunk a bit so Word doesn't choke on giant paragraphs
    for block in content.split("\n\n"):
        p = doc.add_paragraph(block)
        p.paragraph_format.space_after = Pt(6)


def build_report():
    run_dir = RUN_DIR
    assert run_dir.exists(), f"Run folder not found: {run_dir}"

    metrics = _safe_load_json(run_dir / METRICS_JSON, {})
    derived = _compute_derived(metrics)

    # --- start document ---
    doc = Document()

    # default font
    style = doc.styles["Normal"]
    style.font.name = "Calibri"
    style._element.rPr.rFonts.set(qn('w:eastAsia'), 'Calibri')
    style.font.size = Pt(11)

    # Title page
    title = doc.add_heading("Sepsis Diagnostic RL: Final Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph().alignment = WD_ALIGN_PARAGRAPH.CENTER

    meta = doc.add_paragraph(
        f"Run: {run_dir.name}   |   Generated: {datetime.datetime.now().strftime('%Y-%m-%d %H:%M')}"
    )
    meta.alignment = WD_ALIGN_PARAGRAPH.CENTER
    doc.add_paragraph()

    # Executive Summary
    _insert_heading(doc, "Executive Summary", level=1)
    summary = doc.add_paragraph()
    summary.add_run(
        "This document summarizes the final performance of the trained reinforcement-learning (RL) "
        "agent for sepsis mortality prediction and cost-aware test selection. "
        "The results below are taken directly from the evaluation artifacts in "
        f"{run_dir} and presented in a paper-style format."
    )

    doc.add_paragraph()

    # Headline metrics table (single run — no comparisons)
    _insert_heading(doc, "Headline Metrics", level=1)
    rows = [
        ("Diagnostic Accuracy", _fmt_pct(metrics.get("acc"))),
        ("Precision (PPV)", _fmt_float(metrics.get("precision"), 3)),
        ("Recall (Sensitivity)", _fmt_float(metrics.get("recall"), 3)),
        ("Specificity", _fmt_float(metrics.get("specificity"), 3)),
        ("F1-score", _fmt_float(metrics.get("f1"), 3)),
        ("AUROC", _fmt_float(metrics.get("AUC"), 3)),
        ("AUPRC", _fmt_float(metrics.get("AUPRC"), 3)),
        ("Avg Financial Cost / Patient", _fmt_money(metrics.get("avg_cost"))),
        ("Avg Number of Tests Ordered", _fmt_float(metrics.get("avg_steps"), 2)),
        ("Episodes Evaluated", metrics.get("n_episodes", "—")),
        ("Threshold Used", _fmt_float(metrics.get("threshold_used"), 3)),
        ("Threshold Source", metrics.get("threshold_source", "—")),
    ]
    _insert_kv_table(doc, rows)

    # Derived metrics
    _insert_heading(doc, "Derived Metrics", level=2)
    rows2 = [
        ("Prevalence (positives in set)", _fmt_pct(derived["prevalence"])),
        ("NPV", _fmt_float(derived["npv"], 3)),
        ("PPV", _fmt_float(derived["ppv"], 3)),
        ("F1 per Dollar (F1/$)", _fmt_float(derived["f1_per_dollar"], 6)),
        ("TN / FP / FN / TP", f'{metrics.get("tn","—")} / {metrics.get("fp","—")} / {metrics.get("fn","—")} / {metrics.get("tp","—")}'),
    ]
    _insert_kv_table(doc, rows2)

    # Strategy summary (panel usage)
    _insert_heading(doc, "Agent Testing Strategy", level=1)
    doc.add_paragraph(
        "The policy selects laboratory panels adaptively until it chooses to diagnose. "
        "Panel usage below reflects how often each named panel was ordered at least once per episode."
    )
    _insert_image_if_exists(doc, run_dir / IMG_NAMES["Panels"], "Figure: Panel usage (CBC, CMP, aPTT, ABG)")

    # Performance curves and confusion matrix
    _insert_heading(doc, "Performance Curves & Confusion Matrix", level=1)
    _insert_image_if_exists(doc, run_dir / IMG_NAMES["ROC"], "Figure: ROC Curve (AUC)")
    _insert_image_if_exists(doc, run_dir / IMG_NAMES["PR"], "Figure: Precision–Recall Curve (AP)")
    _insert_image_if_exists(doc, run_dir / IMG_NAMES["Confusion"], "Figure: Confusion Matrix (Expired = 1)")

    # Notes
    _insert_heading(doc, "Notes", level=1)
    doc.add_paragraph("• No calibration curves included (per request).")
    doc.add_paragraph("• This report contains only this run’s results; no baseline or cross-run comparison.")
    doc.add_paragraph("• Threshold source is recorded as reported by the evaluator; interpretation is unchanged.")

    # Case studies (optional)
    _append_case_studies(doc, run_dir / CASE_STUDIES)

    # Save
    OUT_DOC.parent.mkdir(parents=True, exist_ok=True)
    doc.save(str(OUT_DOC))
    print(f"[ok] Wrote: {OUT_DOC}")


if __name__ == "__main__":
    build_report()
